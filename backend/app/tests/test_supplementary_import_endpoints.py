"""Integration tests for supplementary import endpoints."""

import io

import pytest
import pytest_asyncio
from httpx import AsyncClient
from openpyxl import Workbook
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.security import require_admin, require_college
from app.main import app
from app.models.college_review import CollegeRanking
from app.models.scholarship import (
    ScholarshipConfiguration,
    ScholarshipSubTypeConfig,
    ScholarshipType,
    SubTypeSelectionMode,
)
from app.models.user import AdminScholarship, User, UserRole, UserType


def _build_xlsx_bytes() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.cell(row=1, column=1, value="Title")
    headers = [
        "NO.",
        "學院初審會議之學院排序",
        "申請獎學金類別",
        "學院",
        "系所",
        "年級",
        "是否為逕博學生",
        "學生中文姓名",
        "學生英文姓名",
        "國籍",
        "性別",
        "註冊入學日期",
        "學號",
        "學生身分證字號",
        "學生匯款帳號",
        "學生E-mail",
        "學生通訊地址",
        "指導教授姓名",
    ]
    for col_idx, h in enumerate(headers, start=1):
        ws.cell(row=2, column=col_idx, value=h)
    row = [
        1,
        1,
        "國科會博士生研究獎學金",
        "工學院",
        "電機系",
        2,
        "否",
        "王小明",
        "Wang",
        "台灣",
        "男",
        "113.9.1",
        "310460099",
        "A123456789",
        "12345678",
        "test99@nycu.edu.tw",
        "新竹市",
        "指導教授A",
    ]
    for col_idx, val in enumerate(row, start=1):
        ws.cell(row=3, column=col_idx, value=val)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest_asyncio.fixture
async def admin_user(db: AsyncSession) -> User:
    user = User(
        nycu_id="admin001",
        name="Admin",
        email="admin@nycu.edu.tw",
        user_type=UserType.employee,
        role=UserRole.admin,
    )
    db.add(user)
    await db.flush()
    return user


@pytest_asyncio.fixture
async def college_user(db: AsyncSession) -> User:
    user = User(
        nycu_id="col001",
        name="College",
        email="col@nycu.edu.tw",
        user_type=UserType.employee,
        role=UserRole.college,
        college_code="A",
    )
    db.add(user)
    await db.flush()
    return user


@pytest_asyncio.fixture
async def scholarship(db: AsyncSession) -> ScholarshipType:
    s = ScholarshipType(
        code="phd_supp_test",
        name="Test Supp PhD",
        sub_type_selection_mode=SubTypeSelectionMode.single,
        status="active",
    )
    db.add(s)
    await db.flush()
    return s


@pytest_asyncio.fixture
async def configuration(db: AsyncSession, scholarship: ScholarshipType, admin_user: User) -> ScholarshipConfiguration:
    cfg = ScholarshipConfiguration(
        scholarship_type_id=scholarship.id,
        academic_year=114,
        semester=None,  # yearly
        config_name="Test PhD 114",
        config_code="test-phd-114",
        amount=40000,
        is_active=True,
        allow_supplementary_import=False,
    )
    db.add(cfg)
    await db.flush()

    # Grant admin permission to manage this scholarship type
    db.add(AdminScholarship(admin_id=admin_user.id, scholarship_id=scholarship.id))
    await db.flush()
    return cfg


@pytest_asyncio.fixture
async def ranking(
    db: AsyncSession,
    college_user: User,
    scholarship: ScholarshipType,
    configuration: ScholarshipConfiguration,
) -> CollegeRanking:
    r = CollegeRanking(
        scholarship_type_id=scholarship.id,
        sub_type_code="nstc",
        academic_year=114,
        ranking_name="Test",
        created_by=college_user.id,
        is_finalized=True,
        distribution_executed=True,
    )
    db.add(r)
    await db.flush()
    return r


@pytest.mark.asyncio
class TestAdminConfigToggle:
    async def test_admin_can_enable_supplementary_import(
        self,
        client: AsyncClient,
        db: AsyncSession,
        admin_user: User,
        configuration: ScholarshipConfiguration,
    ):
        app.dependency_overrides[require_admin] = lambda: admin_user
        try:
            resp = await client.patch(
                f"/api/v1/scholarship-configurations/configurations/{configuration.id}/supplementary-import",
                json={"allow": True},
            )
        finally:
            app.dependency_overrides.pop(require_admin, None)
        assert resp.status_code == 200, resp.text
        data = resp.json()
        assert data["data"]["allow_supplementary_import"] is True

    async def test_returns_404_for_unknown_configuration(self, client: AsyncClient, db: AsyncSession, admin_user: User):
        app.dependency_overrides[require_admin] = lambda: admin_user
        try:
            resp = await client.patch(
                "/api/v1/scholarship-configurations/configurations/999999/supplementary-import",
                json={"allow": True},
            )
        finally:
            app.dependency_overrides.pop(require_admin, None)
        assert resp.status_code == 404


@pytest.mark.asyncio
class TestSupplementaryImportEndpoint:
    async def test_returns_403_when_flag_is_off(
        self,
        client: AsyncClient,
        db: AsyncSession,
        college_user: User,
        ranking: CollegeRanking,
    ):
        app.dependency_overrides[require_college] = lambda: college_user
        try:
            xlsx_bytes = _build_xlsx_bytes()
            resp = await client.post(
                f"/api/v1/college-review/rankings/{ranking.id}/supplementary-import",
                files={
                    "file": (
                        "test.xlsx",
                        xlsx_bytes,
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    )
                },
            )
        finally:
            app.dependency_overrides.pop(require_college, None)
        assert resp.status_code == 403

    async def test_returns_404_for_unknown_ranking(self, client: AsyncClient, db: AsyncSession, college_user: User):
        app.dependency_overrides[require_college] = lambda: college_user
        try:
            xlsx_bytes = _build_xlsx_bytes()
            resp = await client.post(
                "/api/v1/college-review/rankings/999999/supplementary-import",
                files={
                    "file": (
                        "test.xlsx",
                        xlsx_bytes,
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    )
                },
            )
        finally:
            app.dependency_overrides.pop(require_college, None)
        assert resp.status_code == 404

    async def test_rejects_student_from_other_college(
        self,
        client: AsyncClient,
        db: AsyncSession,
        college_user: User,  # college_code = "A"
        ranking: CollegeRanking,
        configuration: ScholarshipConfiguration,
        scholarship: ScholarshipType,
        monkeypatch,
    ):
        """Student whose SIS std_academyno doesn't match the ranking's college
        must be rejected with 422 and the offending student ID listed."""
        # Enable the supplementary-import flag for the configuration
        configuration.allow_supplementary_import = True
        # Seed sub-type config so parse_excel can resolve the Excel label
        db.add(
            ScholarshipSubTypeConfig(
                scholarship_type_id=scholarship.id,
                sub_type_code="nstc",
                name="國科會博士生研究獎學金",
                is_active=True,
            )
        )
        await db.commit()

        # Bypass SIS: return a student whose std_academyno is "EE" (not "A")
        async def fake_fetch(self, student_ids, **kwargs):
            return (
                {sid: {"std_academyno": "EE", "std_cname": "他院學生"} for sid in student_ids},
                [],
            )

        from app.services.supplementary_import_service import SupplementaryImportService

        monkeypatch.setattr(SupplementaryImportService, "fetch_student_data_bulk", fake_fetch)

        app.dependency_overrides[require_college] = lambda: college_user
        try:
            resp = await client.post(
                f"/api/v1/college-review/rankings/{ranking.id}/supplementary-import",
                files={
                    "file": (
                        "test.xlsx",
                        _build_xlsx_bytes(),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    )
                },
            )
        finally:
            app.dependency_overrides.pop(require_college, None)

        assert resp.status_code == 422, resp.text
        body = resp.json()
        # ApiResponse wrapper format
        message = body.get("message") or body.get("detail") or ""
        assert "不屬於本學院" in message
        assert "310460099" in message  # student_id from _build_xlsx_bytes
        assert "EE" in message  # surfaced mismatched college code
