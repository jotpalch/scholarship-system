"""College ranking Excel export service.

Pure rendering logic — receives prepared data structures, returns xlsx bytes.
The endpoint layer is responsible for loading rows, dynamic field configs,
and sub-type labels from the database.
"""

from __future__ import annotations

import io
import math
from dataclasses import dataclass
from typing import Any, Dict, List, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

STATIC_HEADERS: List[str] = [
    "NO.",
    "學院初審會議之學院排序",
    "申請獎學金類別",
    "學院",
    "系所",
    "年級",
    "是否為逕博學生",
    "學生中文姓名",
    "學生英文姓名",
    "國籍",
    "性別",
    "註冊入學日期",
    "學號",
    "學生身分證字號",
    "學生匯款帳號",
    "學生E-mail",
    "學生通訊地址",
    "指導教授姓名",
]


# std_enrolltype codes that indicate 逕讀博士 (direct-track PhD)
DIRECT_PHD_ENROLLTYPE_CODES = {8, 9, 10, 11}


@dataclass(frozen=True)
class DynamicFieldSpec:
    field_name: str
    field_label: str
    export_column_label: Optional[str]
    display_order: int


@dataclass
class ExportRow:
    """One ranked application's data."""

    rank_position: Optional[int]
    application: Any  # Duck-typed: needs sub_type_preferences, sub_scholarship_type, student_data, submitted_form_data
    bank_account: Optional[str] = None  # 郵局帳號 (from user_profiles.account_number)
    advisor_names: Optional[str] = None  # 指導教授姓名 (comma-joined if multiple advisors)


class CollegeRankingExportService:
    """Builds 學生資料彙整表 workbooks."""

    def build_workbook(
        self,
        *,
        rows: List[ExportRow],
        dynamic_fields: List[DynamicFieldSpec],
        sub_type_labels: Dict[str, str],
        title: str,
        sheet_name: str,
    ) -> bytes:
        wb = Workbook()
        ws = wb.active
        ws.title = sheet_name

        sorted_dynamic = sorted(dynamic_fields, key=lambda f: (f.display_order, f.field_name))
        headers = STATIC_HEADERS + [(f.export_column_label or f.field_label) for f in sorted_dynamic]
        total_cols = len(headers)

        # Row 1: title (merged across all columns)
        ws.cell(row=1, column=1, value=title)
        if total_cols > 1:
            ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=total_cols)
        title_cell = ws.cell(row=1, column=1)
        title_cell.font = Font(bold=True, size=14)
        title_cell.alignment = Alignment(horizontal="center", vertical="center")

        # Row 2: header
        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=2, column=col_idx, value=header)
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.fill = PatternFill("solid", fgColor="DDDDDD")

        # Data rows
        for idx, row in enumerate(rows, start=1):
            excel_row = idx + 2  # +2 because rows 1-2 are title/header
            self._write_static_cells(ws, excel_row, idx, row, sub_type_labels)
            self._write_dynamic_cells(ws, excel_row, row, sorted_dynamic)

        max_row = len(rows) + 2
        self._apply_borders(ws, max_row=max_row, max_col=total_cols)
        self._apply_column_widths(ws, headers)
        ws.freeze_panes = "A3"

        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    # -------- Static column rendering --------

    def _write_static_cells(
        self,
        ws,
        excel_row: int,
        row_index: int,
        row: ExportRow,
        sub_type_labels: Dict[str, str],
    ) -> None:
        sd = getattr(row.application, "student_data", None) or {}

        ws.cell(row=excel_row, column=1, value=row_index)
        ws.cell(
            row=excel_row,
            column=2,
            value=(row.rank_position if row.rank_position is not None else ""),
        )
        ws.cell(
            row=excel_row,
            column=3,
            value=self._render_scholarship_type(row.application, sub_type_labels),
        )
        ws.cell(row=excel_row, column=4, value=self._safe_str(sd.get("trm_academyname")))
        ws.cell(row=excel_row, column=5, value=self._safe_str(sd.get("trm_depname")))
        ws.cell(row=excel_row, column=6, value=self._compute_grade(sd.get("trm_termcount")))
        ws.cell(row=excel_row, column=7, value=self._render_direct_phd(sd.get("std_enrolltype")))
        ws.cell(row=excel_row, column=8, value=self._safe_str(sd.get("std_cname")))
        ws.cell(row=excel_row, column=9, value=self._safe_str(sd.get("std_ename")))
        ws.cell(row=excel_row, column=10, value=self._safe_str(sd.get("std_nation")))
        ws.cell(row=excel_row, column=11, value=self._render_gender(sd.get("std_sex")))
        ws.cell(
            row=excel_row,
            column=12,
            value=self._render_enrollment_date(sd.get("std_enrollyear"), sd.get("std_enrollterm")),
        )
        ws.cell(row=excel_row, column=13, value=self._safe_str(sd.get("std_stdcode")))
        ws.cell(row=excel_row, column=14, value=self._safe_str(sd.get("std_pid")))
        ws.cell(row=excel_row, column=15, value=self._safe_str(row.bank_account))
        ws.cell(row=excel_row, column=16, value=self._safe_str(sd.get("com_email")))
        ws.cell(row=excel_row, column=17, value=self._safe_str(sd.get("com_commadd")))
        ws.cell(row=excel_row, column=18, value=self._safe_str(row.advisor_names))

    def _safe_str(self, value: Any) -> str:
        if value is None:
            return ""
        return str(value)

    def _render_gender(self, std_sex: Any) -> str:
        if std_sex == 1:
            return "男"
        if std_sex == 2:
            return "女"
        return ""

    def _render_direct_phd(self, std_enrolltype: Any) -> str:
        try:
            code = int(std_enrolltype)
        except (TypeError, ValueError):
            return ""
        return "是" if code in DIRECT_PHD_ENROLLTYPE_CODES else "否"

    def _compute_grade(self, trm_termcount: Any):
        try:
            tc = int(trm_termcount)
        except (TypeError, ValueError):
            return ""
        if tc <= 0:
            return ""
        return math.ceil(tc / 2)

    def _render_enrollment_date(self, year: Any, term: Any) -> str:
        if year in (None, ""):
            return ""
        try:
            year_int = int(year)
        except (TypeError, ValueError):
            return ""
        month = 9 if term in (None, "", 1, "1") else 2
        return f"{year_int}.{month}.1"

    def _render_scholarship_type(self, application: Any, sub_type_labels: Dict[str, str]) -> str:
        prefs: Optional[List[str]] = getattr(application, "sub_type_preferences", None)
        if prefs:
            labels = [sub_type_labels.get(code, code) for code in prefs]
            if len(labels) == 1:
                return labels[0]
            return f"{labels[0]}(第一志願)暨{labels[1]}(第二志願)"

        fallback_code = getattr(application, "sub_scholarship_type", None)
        if not fallback_code:
            return ""
        return sub_type_labels.get(fallback_code, fallback_code)

    # -------- Dynamic column rendering --------

    def _write_dynamic_cells(
        self,
        ws,
        excel_row: int,
        row: ExportRow,
        dynamic_fields: List[DynamicFieldSpec],
    ) -> None:
        form_data = getattr(row.application, "submitted_form_data", None) or {}
        fields_map = form_data.get("fields") if isinstance(form_data, dict) else None
        fields_map = fields_map if isinstance(fields_map, dict) else {}

        for offset, spec in enumerate(dynamic_fields):
            col_idx = len(STATIC_HEADERS) + 1 + offset
            value = self._extract_dynamic_value(fields_map, spec.field_name)
            ws.cell(row=excel_row, column=col_idx, value=value)

    def _extract_dynamic_value(self, fields_map: Dict[str, Any], field_name: str) -> str:
        entry = fields_map.get(field_name)
        if not isinstance(entry, dict):
            return ""
        raw = entry.get("value")
        if raw is None or raw == "":
            return ""
        return str(raw)

    # -------- Formatting --------

    def _apply_borders(self, ws, *, max_row: int, max_col: int) -> None:
        thin = Side(style="thin", color="999999")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        for r in range(2, max_row + 1):
            for c in range(1, max_col + 1):
                ws.cell(row=r, column=c).border = border

    def _apply_column_widths(self, ws, headers: List[str]) -> None:
        for idx, header in enumerate(headers, start=1):
            text_len = max(len(str(header)), 6)
            width = min(max(text_len + 4, 10), 30)
            ws.column_dimensions[get_column_letter(idx)].width = width
